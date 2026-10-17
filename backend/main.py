
import json
import re
from io import BytesIO
from threading import Lock

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel, Field

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from test_workers import graph, llm


app = FastAPI(
    title="Blog Writing Agent API",
    description="Generate and revise blogs using Gemini, LangGraph and Tavily.",
    version="1.4.0",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://127.0.0.1:5500",
        "http://localhost:5500",
        "http://127.0.0.1:8000",
        "http://localhost:8000",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["Content-Type"],
)


class BlogRequest(BaseModel):
    topic: str = Field(min_length=3, max_length=300)
    blog_length: str = Field(
        default="medium",
        pattern="^(short|medium|long|detailed)$",
    )
    writing_tone: str = Field(
        default="professional",
        pattern="^(professional|conversational|friendly|academic|persuasive)$",
    )
    target_audience: str = Field(
        default="general",
        pattern="^(general|beginners|students|professionals|technical)$",
    )


class BlogRevisionRequest(BaseModel):
    topic: str = Field(default="Blog topic", min_length=1, max_length=300)
    blog_markdown: str = Field(min_length=20, max_length=30000)
    instructions: str = Field(min_length=3, max_length=1500)


class BlogExportRequest(BaseModel):
    topic: str = Field(
        default="InkFlow AI Blog",
        min_length=1,
        max_length=300,
    )
    blog_markdown: str = Field(min_length=20, max_length=30000)


generation_lock = Lock()


def initial_state(
    topic: str,
    blog_length: str = "medium",
    writing_tone: str = "professional",
    target_audience: str = "general",
):
    return {
        "topic": topic,
        "blog_length": blog_length,
        "writing_tone": writing_tone,
        "target_audience": target_audience,
        "needs_research": "",
        "research": "",
        "sources": [],
        "sections": [],
        "completed_sections": [],
        "final_blog": "",
    }


@app.get("/api/health")
def health_check():
    return {
        "status": "ok",
        "message": "Blog Writing Agent backend is running!",
    }


@app.post("/api/generate")
def generate_blog(request: BlogRequest):
    topic = request.topic.strip()

    if len(topic) < 3:
        raise HTTPException(
            status_code=400,
            detail="Please enter a topic with at least 3 characters.",
        )

    state = initial_state(
        topic=topic,
        blog_length=request.blog_length,
        writing_tone=request.writing_tone,
        target_audience=request.target_audience,
    )

    try:
        with generation_lock:
            result = graph.invoke(state)

        return {
            "success": True,
            "topic": topic,
            "blog_length": request.blog_length,
            "writing_tone": request.writing_tone,
            "target_audience": request.target_audience,
            "blog_markdown": result["final_blog"],
            "sources": result.get("sources", []),
        }

    except Exception as error:
        print(f"Blog generation failed: {error}")
        raise HTTPException(
            status_code=500,
            detail=(
                "Blog generation failed. "
                "Check the backend terminal for details."
            ),
        )


# ============================================================
# WORD DOCX EXPORT
# Creates a standard DOCX using python-docx on the backend.
# ============================================================

@app.post("/api/export/docx")
def export_blog_docx(request: BlogExportRequest):
    blog = request.blog_markdown.strip()
    topic = request.topic.strip()

    if len(blog) < 20:
        raise HTTPException(
            status_code=400,
            detail="The blog is too short to export.",
        )

    try:
        doc = Document()

        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        normal = doc.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.15

        for style_name, size in (
            ("Title", 22),
            ("Heading 1", 16),
            ("Heading 2", 13),
            ("Heading 3", 12),
        ):
            style = doc.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor(35, 42, 75)

        doc.core_properties.title = topic or "InkFlow AI Blog"
        doc.core_properties.subject = "Blog exported from InkFlow AI"
        doc.core_properties.author = "InkFlow AI"

        for raw_line in blog.replace("\r\n", "\n").split("\n"):
            line = raw_line.strip()

            if not line:
                continue

            # Remove Markdown code-fence markers.
            if line.startswith("```"):
                continue

            # Markdown headings.
            heading = re.match(r"^(#{1,3})\s+(.+)$", line)

            if heading:
                level = len(heading.group(1))
                doc.add_paragraph(
                    heading.group(2),
                    style=f"Heading {level}",
                )
                continue

            # Recognize common section headings without # markers.
            if line.lower() in {
                "introduction",
                "conclusion",
                "sources",
            }:
                doc.add_paragraph(line, style="Heading 1")
                continue

            # Lists.
            bullet = re.match(r"^\s*[-*•]\s+(.+)$", line)
            numbered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)

            if bullet:
                paragraph = doc.add_paragraph(style="List Bullet")
                content = bullet.group(1)

            elif numbered:
                paragraph = doc.add_paragraph(style="List Number")
                content = numbered.group(1)

            else:
                paragraph = doc.add_paragraph()
                content = line

            # Convert Markdown links to readable text with URLs.
            content = re.sub(
                r"\[([^\]]+)\]\((https?://[^\s)]+)\)",
                r"\1 (\2)",
                content,
            )

            content = re.sub(
                r"!\[([^\]]*)\]\([^)]+\)",
                r"\1",
                content,
            )

            # Remove Markdown bold, italic and inline-code markers.
            content = re.sub(
                r"(\*\*|__)(.*?)\1",
                r"\2",
                content,
            )
            content = re.sub(
                r"(\*|_)(.*?)\1",
                r"\2",
                content,
            )
            content = re.sub(
                r"`([^`]+)`",
                r"\1",
                content,
            )

            paragraph.add_run(content)

        # Write a fresh Word document into memory.
        buffer = BytesIO()
        doc.save(buffer)
        file_bytes = buffer.getvalue()

        safe_topic = re.sub(
            r"[^a-zA-Z0-9_-]+",
            "-",
            topic,
        ).strip("-_")[:60] or "inkflow-blog"

        filename = f"{safe_topic}.docx"

        return Response(
            content=file_bytes,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{filename}"'
                ),
                "Content-Length": str(len(file_bytes)),
                "Cache-Control": "no-store",
            },
        )

    except Exception as error:
        print(f"DOCX export failed: {error}")
        raise HTTPException(
            status_code=500,
            detail=(
                "Word export failed. "
                "Check the backend terminal for details."
            ),
        )


# ============================================================
# BLOG REGENERATION
# ============================================================

@app.post("/api/regenerate")
def regenerate_blog(request: BlogRevisionRequest):
    blog = request.blog_markdown.strip()
    instructions = request.instructions.strip()
    topic = request.topic.strip()

    if len(blog) < 20:
        raise HTTPException(
            status_code=400,
            detail="The blog is too short to revise. Generate a blog first.",
        )

    if len(instructions) < 3:
        raise HTTPException(
            status_code=400,
            detail="Please describe the changes you want.",
        )

    prompt = f"""
You are an expert blog editor. Revise the existing blog using the user's
instructions. Return the complete revised blog in Markdown, not an explanation
of your changes.

BLOG TOPIC:
{topic}

USER'S REVISION INSTRUCTIONS:
{instructions}

CURRENT BLOG:
{blog}

Rules:
- Follow the user's revision instructions as closely as possible.
- Return the entire revised blog, not only the changed paragraphs.
- Preserve the topic, useful information, headings, and Markdown structure
  unless the user asks to change them.
- Preserve the existing Sources section and its URLs exactly where possible.
- Do not invent citations, URLs, statistics, quotations, or factual claims.
- If asked to add factual information not supported by the current blog,
  qualify it cautiously rather than pretending it has been verified.
- Do not include commentary such as "Here is your revised blog".
- Return only the revised blog in Markdown.
"""

    try:
        with generation_lock:
            response = llm.invoke(prompt)

        content = response.content

        if isinstance(content, list):
            content = "\n".join(
                item["text"]
                for item in content
                if isinstance(item, dict)
                and item.get("type") == "text"
            )

        revised_blog = str(content).strip()

        if not revised_blog:
            raise RuntimeError("Gemini returned an empty revised blog.")

        return {
            "success": True,
            "topic": topic,
            "blog_markdown": revised_blog,
        }

    except Exception as error:
        print(f"Blog revision failed: {error}")
        raise HTTPException(
            status_code=500,
            detail=(
                "Blog revision failed. "
                "Check the backend terminal for details."
            ),
        )


# ============================================================
# STREAMING BLOG GENERATION
# ============================================================

def event_stream(
    topic: str,
    blog_length: str = "medium",
    writing_tone: str = "professional",
    target_audience: str = "general",
):
    def send_event(event_name: str, **data):
        payload = {"event": event_name, **data}
        return f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

    yield send_event(
        "status",
        stage="starting",
        message="Preparing your blog request...",
    )

    state = initial_state(
        topic=topic,
        blog_length=blog_length,
        writing_tone=writing_tone,
        target_audience=target_audience,
    )

    final_state = state.copy()

    messages = {
        "router": (
            "research",
            "Checking whether web research is needed...",
        ),
        "research": (
            "research",
            "Web research completed.",
        ),
        "skip_research": (
            "research",
            "Web research is not required for this topic.",
        ),
        "planner": (
            "planning",
            "Blog outline created. Preparing sections...",
        ),
        "reducer": (
            "assembly",
            "Assembling the introduction and conclusion...",
        ),
        "source_verifier": (
            "assembly",
            "Reviewing factual claims against available sources...",
        ),
        "proofreader": (
            "proofreading",
            "Proofreading the completed blog...",
        ),
    }

    try:
        yield send_event(
            "status",
            stage="working",
            message="Starting the AI workflow...",
        )

        with generation_lock:
            for update in graph.stream(
                state,
                stream_mode="updates",
            ):
                for node_name, node_update in update.items():
                    if not isinstance(node_update, dict):
                        continue

                    for key, value in node_update.items():
                        if key == "completed_sections":
                            final_state[key] = (
                                final_state.get(key, []) + value
                            )
                        else:
                            final_state[key] = value

                    if node_name == "worker":
                        completed = len(
                            final_state.get("completed_sections", [])
                        )
                        total = len(
                            final_state.get("sections", [])
                        ) or 5

                        yield send_event(
                            "status",
                            stage="writing",
                            message=(
                                f"Writing blog sections: "
                                f"{completed}/{total} completed."
                            ),
                            completed=completed,
                            total=total,
                        )

                    elif node_name in messages:
                        stage, message = messages[node_name]

                        yield send_event(
                            "status",
                            stage=stage,
                            message=message,
                        )

        blog = final_state.get("final_blog", "")

        if not blog:
            raise RuntimeError(
                "The workflow finished without producing a blog."
            )

        yield send_event(
            "status",
            stage="complete",
            message="Blog generation completed successfully.",
        )

        yield send_event(
            "result",
            success=True,
            topic=topic,
            blog_length=blog_length,
            writing_tone=writing_tone,
            target_audience=target_audience,
            blog_markdown=blog,
            sources=final_state.get("sources", []),
        )

    except Exception as error:
        print(f"Streaming blog generation failed: {error}")

        yield send_event(
            "error",
            message=(
                "Blog generation failed. "
                "Check the backend terminal for details."
            ),
        )


@app.post("/api/generate/stream")
def generate_blog_stream(request: BlogRequest):
    topic = request.topic.strip()

    if len(topic) < 3:
        raise HTTPException(
            status_code=400,
            detail="Please enter a topic with at least 3 characters.",
        )

    return StreamingResponse(
        event_stream(
            topic=topic,
            blog_length=request.blog_length,
            writing_tone=request.writing_tone,
            target_audience=request.target_audience,
        ),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        "backend.main:app",
        host="127.0.0.1",
        port=8000,
        reload=True,
    )
